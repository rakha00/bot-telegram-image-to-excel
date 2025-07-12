

import os
import json
import uuid
from datetime import datetime
import logging

from flask import Flask, request, jsonify, render_template, send_from_directory
from werkzeug.utils import secure_filename
import docx
import pdfplumber

import gemini_vision_extractor

# Setup logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Initialize Flask App
app = Flask(__name__)

# Configuration
UPLOAD_FOLDER = 'temp_files'
OUTPUT_FOLDER = 'output'
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'pdf', 'docx'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['OUTPUT_FOLDER'] = OUTPUT_FOLDER

# Ensure directories exist
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def fix_empty_key(json_data, new_key="Akun"):
    if not json_data:
        return json_data
    old_key = ""
    if isinstance(json_data, list) and len(json_data) > 0 and old_key in json_data[0]:
        for obj in json_data:
            if isinstance(obj, dict) and old_key in obj:
                obj[new_key] = obj.pop(old_key)
    return json_data

def pdf_to_json(pdf_path):
    logger.info(f"Opening PDF: {pdf_path}")
    with pdfplumber.open(pdf_path) as pdf:
        first_page = pdf.pages[0]
        tables = first_page.extract_tables()
        if not tables:
            logger.warning(f"No tables found in PDF: {pdf_path}")
            return []
        table = tables[0]
        headers = table[0]
        data = []
        for row in table[1:]:
            obj = {}
            for i, cell in enumerate(row):
                key = headers[i] if i < len(headers) else f"col_{i+1}"
                obj[key] = cell if cell not in [None, ""] else None
            data.append(obj)
        logger.info(f"Successfully extracted {len(data)} rows from PDF.")
        return data

def docx_to_json(docx_path):
    logger.info(f"Opening DOCX: {docx_path}")
    doc = docx.Document(docx_path)
    if not doc.tables:
        logger.warning(f"No tables found in DOCX: {docx_path}")
        return []
    table = doc.tables[0]
    rows = list(table.rows)
    headers = [cell.text.strip() for cell in rows[0].cells]
    data = []
    for row in rows[1:]:
        obj = {}
        for i, cell in enumerate(row.cells):
            key = headers[i] if i < len(headers) else f"col_{i+1}"
            value = cell.text.strip()
            obj[key] = value if value else None
        data.append(obj)
    logger.info(f"Successfully extracted {len(data)} rows from DOCX.")
    return data

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file part"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400
        
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_ext = filename.rsplit('.', 1)[1].lower()
        
        # Save the uploaded file temporarily
        temp_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(temp_path)
        logger.info(f"File saved temporarily to {temp_path}")

        data = None
        try:
            if file_ext in {'png', 'jpg', 'jpeg', 'gif'}:
                # Process image with Gemini
                logger.info(f"Processing image file: {temp_path}")
                json_result = ""
                # The gemini extractor is async, but we'll call it in a blocking way for this web context
                # Note: This is a simplified approach. For production, you might run this in a separate thread.
                import asyncio
                json_result = asyncio.run(gemini_vision_extractor.get_json_output(temp_path))

                if json_result.strip().startswith("```"):
                    json_result = json_result.strip().lstrip("`json").lstrip("`").strip()
                    if json_result.endswith("```"):
                        json_result = json_result[:json_result.rfind("```")].strip()
                
                data = json.loads(json_result)

            elif file_ext == 'pdf':
                logger.info(f"Processing PDF file: {temp_path}")
                data = pdf_to_json(temp_path)

            elif file_ext == 'docx':
                logger.info(f"Processing DOCX file: {temp_path}")
                data = docx_to_json(temp_path)

            if data is not None:
                # Clean up the data
                data = fix_empty_key(data, new_key="Akun")
                
                # Save the final JSON to the output folder
                base_filename = os.path.splitext(filename)[0]
                timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
                unique_id = uuid.uuid4().hex[:8]
                output_filename = f"{base_filename}_{timestamp}_{unique_id}.json"
                output_path = os.path.join(app.config['OUTPUT_FOLDER'], output_filename)
                
                with open(output_path, 'w', encoding='utf-8') as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)
                
                logger.info(f"Successfully processed file. Result saved to {output_path}")
                return jsonify({"success": True, "data": data, "download_url": f"/download/{output_filename}"})
            else:
                return jsonify({"error": "Could not extract any data from the file."}), 500

        except Exception as e:
            logger.error(f"An error occurred: {e}", exc_info=True)
            return jsonify({"error": str(e)}), 500
        finally:
            # Clean up the temporary file
            if os.path.exists(temp_path):
                os.remove(temp_path)
                logger.info(f"Removed temporary file: {temp_path}")
    
    return jsonify({"error": "File type not allowed"}), 400

@app.route('/download/<filename>')
def download_file(filename):
    return send_from_directory(app.config["OUTPUT_FOLDER"], filename, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5001)

