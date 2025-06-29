"""
Modul untuk mengekstrak tabel dari gambar menggunakan Google Gemini Vision API.
"""
import os
import google.generativeai as genai
import PIL.Image
import asyncio

def configure_gemini():
    """Konfigurasi Gemini API dengan kunci dari environment variables."""
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("GEMINI_API_KEY tidak ditemukan di file .env")
    genai.configure(api_key=api_key)

def generate_gemini_prompt():
    """Membangun prompt yang sangat ketat untuk menghasilkan skrip Python."""
    return """
    **MISI ANDA: UBAH GAMBAR MENJADI SKRIP PYTHON YANG MEREPLIKASI GRID EXCEL DENGAN SEMPURNA.**

    **PERINTAH UTAMA**:
    - **HANYA KODE PYTHON.** Seluruh respons Anda harus berupa kode Python mentah yang dapat dieksekusi. Jangan tambahkan kata lain, penjelasan, atau format markdown.
    - **FOKUS PADA STRUKTUR GRID.** Prioritas utama Anda adalah menempatkan data di sel yang benar.
    - Gunakan `pandas` untuk data dan `openpyxl` untuk `merge_cells`.

    **ATURAN PENTING - JANGAN LAKUKAN INI**:
    - **JANGAN** gabungkan sel kecuali Anda melihat satu sel dengan jelas membentang beberapa baris atau kolom dalam gambar.
    - **JANGAN** mengarang data atau baris yang tidak ada dalam gambar. Jika sel kosong, gunakan `None`.
    - **JANGAN** gunakan `openpyxl.styles` atau mencoba mereplikasi gaya visual apa pun (tebal, perataan, dll.).

    **CONTOH SEDERHANA**:
    import pandas as pd
    from openpyxl import load_workbook

    def create_excel(output_path: str):
        data = [
            ['Header 1', 'Header 2', 'Header 3'],
            ['Data A1', 'Data B1', 'Data C1'],
            ['Data A2', None, 'Data C2']
        ]
        df = pd.DataFrame(data)
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Sheet1', index=False, header=False)
            workbook = writer.book
            worksheet = writer.sheets['Sheet1']
            # worksheet.merge_cells('A1:C1')
            workbook.save(output_path)
    """

async def stream_excel_script(image_path: str, model_name: str = 'gemini-1.5-flash'):
    """
    Menghasilkan skrip Python secara streaming untuk membuat file Excel dari gambar menggunakan Gemini.
    """
    try:
        configure_gemini()
        
        model = genai.GenerativeModel(model_name)
        prompt = generate_gemini_prompt()
        image = PIL.Image.open(image_path)

        response_stream = await model.generate_content_async([prompt, image], stream=True)
        
        async for chunk in response_stream:
            if chunk.text:
                print(chunk.text, end='', flush=True)
                yield chunk.text

    except Exception as e:
        print(f"Error saat streaming dari Gemini: {e}")
        yield ""